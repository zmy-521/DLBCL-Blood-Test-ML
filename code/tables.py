from pathlib import Path
import sys,json
from config import O, F, prepare_run
import pandas as pd,numpy as np
from scipy.stats import mannwhitneyu,chi2_contingency,fisher_exact
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=O/'tables_v2';R=O/'results_v2'
# Candidate order is loaded from public metadata by config.
units={'AGE':'years','ALT':'U/L','AST':'U/L','TP':'g/L','ALB':'g/L','GLB':'g/L','A/G':'ratio','TBIL':'µmol/L','DBIL':'µmol/L','IBIL':'µmol/L','ALP':'U/L','GGT':'U/L','PAB':'mg/L','BUN':'mmol/L','CREA':'µmol/L','UA':'µmol/L','GLU':'mmol/L',**{f:'10⁹/L' for f in ['WBC','NEUT#','LYMPH#','MONO#','EO#','BASO#','PLT']},**{f:'%' for f in ['NEUT%','LYMPH%','MONO%','EO%','BASO%','HCT','RDW-CV','PCT','P-LCR']},'RBC':'10¹²/L','HGB':'g/L','MCV':'fL','MCH':'pg','MCHC':'g/L','RDW-SD':'fL','PDW':'fL','MPV':'fL'}
def ptext(p):return '<0.001' if p<.001 else f'{p:.3f}'
def three_line(filename,title,headers,rows,notes,landscape=False):
 d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.65)
 if landscape:s.page_width=Inches(11.7);s.page_height=Inches(8.3)
 style=d.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(3)
 p=d.add_paragraph();r=p.add_run(title);r.bold=True;r.font.size=Pt(11)
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 widths=[(s.page_width-s.left_margin-s.right_margin)/len(headers)]*len(headers)
 if len(headers)==4:widths=[Inches(2.25),Inches(1.6),Inches(1.6),Inches(.7)]
 for c,x,w in zip(t.rows[0].cells,headers,widths):c.text=str(x);c.width=int(w)
 head=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');head.append(rep)
 for row in rows:
  cs=t.add_row().cells
  for c,x,w in zip(cs,row,widths):c.text=str(x);c.width=int(w)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  for c in row.cells:
   for p in c.paragraphs:
    for r in p.runs:r.font.name='Times New Roman';r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0);r.font.size=Pt(9 if landscape else 10);r.bold=ri==0
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    el=OxmlElement('w:'+edge);on=(edge=='top' and ri==0)or(edge=='bottom' and ri in [0,len(t.rows)-1]);el.set(qn('w:val'),'single' if on else 'nil');el.set(qn('w:sz'),'8' if ri in [0,len(t.rows)-1] else '4');el.set(qn('w:color'),'000000');b.append(el)
   pr.append(b)
 for note in notes:d.add_paragraph(note)
 d.save(T/filename)
def descriptive(center,filename,title):
 d=pd.read_csv(O/f'data_v2/{center}_strict_baseline_anonymous.csv');assert d.group_key.nunique()==len(d)
 a=d[d.group.eq(0)];b=d[d.group.eq(1)];rows=[];data=[]
 for f in ['AGE']+F:
  x=pd.to_numeric(a[f],errors='coerce').dropna();y=pd.to_numeric(b[f],errors='coerce').dropna();p=mannwhitneyu(x,y,alternative='two-sided',method='asymptotic').pvalue
  disp=lambda s:f'{s.median():.2f} ({s.quantile(.25):.2f}, {s.quantile(.75):.2f})'
  row=[f'{f} ({units[f]})',disp(x),disp(y),ptext(p)];rows.append(row);data.append(dict(variable=f,n_control=len(x),n_case=len(y),control_median=x.median(),control_q1=x.quantile(.25),control_q3=x.quantile(.75),case_median=y.median(),case_q1=y.quantile(.25),case_q3=y.quantile(.75),p=p,test='Mann–Whitney U'))
 sex=np.array([[(a.SEX=='M').sum(),(a.SEX=='F').sum()],[(b.SEX=='M').sum(),(b.SEX=='F').sum()]]);chi,p,_,expected=chi2_contingency(sex,correction=False)
 test='Pearson chi-square'
 if expected.min()<5:p=fisher_exact(sex).pvalue;test='Fisher exact'
 rows.insert(1,['Male, n (%)',f'{sex[0,0]} ({100*sex[0,0]/len(a):.1f})',f'{sex[1,0]} ({100*sex[1,0]/len(b):.1f})',ptext(p)])
 data.append(dict(variable='Male',n_control=len(a),n_case=len(b),control_n=int(sex[0,0]),case_n=int(sex[1,0]),p=p,test=test));pd.DataFrame(data).to_csv(R/(Path(filename).stem+'_statistics.csv'),index=False)
 notes=['Continuous variables: median (Q1, Q3). Categorical variables: n (%). Two-sided Mann–Whitney U tests for continuous variables; '+test+' for sex. P values are exploratory and unadjusted for multiplicity.',f'Each person contributes one record. Cases: earliest eligible paired pretreatment panel with reliable sampling chronology ({len(b)} patients). Controls: one panel per person, as formally confirmed by the study investigator.','This strict-baseline table is distinct from the group-aware repeated-measure primary modeling cohort. Laboratory panel pairing requires the same verified patient and a unique chemistry/CBC panel on the same sampling date.','BUN follows the source label (urea/urea nitrogen, mmol/L); assay nomenclature requires laboratory confirmation. Healthy controls do not represent symptomatic disease controls.']
 three_line(filename,title,["Characteristic",f'Healthy controls (n={len(a)})',f'DLBCL (n={len(b)})','P value'],rows,notes)
def performance_tables():
 p=R/'patient_level_performance.csv'
 if not p.exists():return
 d=pd.read_csv(p);spec=json.loads((O/'models_v2/frozen_model_specification.json').read_text(encoding='utf8'))['models']['Final procedure'];needed=set(['LYMPH#','LYMPH%','HGB','HCT','TP']+spec['features'])
 for typ,label in [('models','Table S2. Patient-level performance of nested models'),('single','Table S3. Single-variable comparators')]:
  rows=[]
  for _,r in d.iterrows():
   if typ=='models' and r.model.startswith('Single:'):continue
   if typ=='single' and not (r.model.startswith('Single: ') and r.model[8:] in needed):continue
   ci=lambda x:f'{r[x]:.3f} ({r[x+"_low"]:.3f}–{r[x+"_high"]:.3f})'
   rows.append([r.cohort,r.model,ci('AUROC'),ci('AUPRC'),f'{r.sensitivity:.3f}',f'{r.specificity:.3f}',f'{r.accuracy:.3f}',f'{r.F1:.3f}',f'{r.PPV:.3f}',f'{r.NPV:.3f}'])
  three_line(f'{typ}_performance.docx',label,['Cohort','Model','AUROC (95% CI)','AUPRC (95% CI)','Sens.','Spec.','Accuracy','F1','PPV','NPV'],rows,['Patient-level metrics average probabilities across eligible repeated panels. CIs: 2000 stratified patient bootstrap samples. Internal classification uses training-fold thresholds; external uses development-frozen thresholds. PPV/NPV depend on case/control sampling fractions.','Primary: development 65 cases/430 controls; external 131/248. Strict-baseline sensitivity: development 65/430; external 41/248. Final procedure: nested selection of algorithm, parameters and feature count. OOF evaluates an adaptive procedure, not one identical fixed model across folds.'],True)
 rows=[]
 for _,r in d[~d.model.str.startswith('Single:')].iterrows():rows.append([r.cohort,r.model,f'{r.calibration_intercept:.3f}',f'{r.calibration_slope:.3f}',f'{r.calibration_in_the_large:.3f}',f'{r.brier:.4f}',str(r.calibration_converged)])
 three_line('calibration_results.docx','Table S4. Patient-level calibration',['Cohort','Model','Joint intercept','Slope','CITL intercept','Brier score','Converged'],rows,['Joint intercept and slope are estimated together. Calibration-in-the-large (CITL) fixes slope at 1. Evaluation only: no external recalibration.'],True)
def main():
 prepare_run()
 descriptive('development','Table_1_patient_baseline.docx','Table 1. Development strict-baseline patient characteristics')
 descriptive('external','Table_S1_patient_baseline.docx','Table S1. External strict-baseline patient characteristics')
 performance_tables();print('Tables saved.')

if __name__=='__main__':main()
